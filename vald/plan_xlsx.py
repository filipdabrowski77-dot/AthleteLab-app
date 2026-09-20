"""Eksport planu treningowego do xlsx — arkusz na trening, kolumny
Nr | Ćwiczenie | Week 1..N | Uwagi | Film.

Zasady Filipa: bez wypełnień kolorem i bez conditional formatting,
bold tylko w nagłówkach. Dawki tygodni pokazujemy EFEKTYWNE (pusty
tydzień dziedziczy z poprzedniego — jak w edytorze i w PDF).
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .plan_pdf_adapter import _kap, _txt
from .training import ensure_sessions, week_params

_SEKCJE = (("Prep", "MOVEMENT PREP"),
           ("Plyo & Power", "PLYO & POWER"),
           ("Main", "MAIN STRENGTH"))


def _dawka(it: dict, wk: int) -> str:
    p = week_params(it, wk)
    s_n, r = str(p.get("sets_n") or ""), str(p.get("reps") or "")
    out = f"{s_n} x {r}" if s_n and r else (r or s_n)
    if out and p.get("intent"):
        out += f" @ {p['intent']}"
    return out


def build_plan_xlsx(plan: dict, url_map: dict | None = None) -> bytes:
    yt = url_map or {}
    weeks_n = max(int(plan.get("weeks", 1) or 1), 1)
    wb = Workbook()
    wb.remove(wb.active)
    naglowek = Font(bold=True)
    srodek = Alignment(horizontal="center")

    sessions = ensure_sessions(plan) or []
    for idx, sess in enumerate(sessions):
        litera = chr(ord("A") + idx) if idx < 26 else str(idx + 1)
        ws = wb.create_sheet(f"Trening {litera}")
        ws.append([_txt(plan.get("name", "")), "",
                   _txt(plan.get("athlete", ""))])
        ws["A1"].font = naglowek
        wiersz = 3

        items = sess.get("items") or []
        for sec_key, sec_tytul in _SEKCJE:
            if sec_key == "Main":
                sec_items = [it for it in items
                             if it.get("section") not in
                             ("Prep", "Plyo & Power")]
            else:
                sec_items = [it for it in items
                             if it.get("section") == sec_key]
            sec_items = [it for it in sec_items
                         if (it.get("exercise") or "").strip()]
            if not sec_items:
                continue
            ws.cell(row=wiersz, column=1, value=sec_tytul).font = naglowek
            wiersz += 1
            kol = ["Nr", "Ćwiczenie"] \
                + [f"Week {k}" for k in range(1, weeks_n + 1)] \
                + ["Uwagi", "Film"]
            for c, v in enumerate(kol, start=1):
                kom = ws.cell(row=wiersz, column=c, value=v)
                kom.font = naglowek
                if 3 <= c <= 2 + weeks_n:
                    kom.alignment = srodek
            wiersz += 1
            for it in sec_items:
                nazwa = _kap(_txt(it.get("exercise") or ""))
                ws.cell(row=wiersz, column=1,
                        value=it.get("slot", "") or "")
                ws.cell(row=wiersz, column=2, value=nazwa)
                for k in range(1, weeks_n + 1):
                    kom = ws.cell(row=wiersz, column=2 + k,
                                  value=_dawka(it, k) or "—")
                    kom.alignment = srodek
                ws.cell(row=wiersz, column=3 + weeks_n,
                        value=_txt(it.get("note") or ""))
                url = yt.get((it.get("exercise") or "").strip().lower())
                if url:
                    kom = ws.cell(row=wiersz, column=4 + weeks_n,
                                  value="FILM")
                    kom.hyperlink = url
                    kom.font = Font(color="0563C1", underline="single")
                wiersz += 1
            wiersz += 1

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 34
        for k in range(1, weeks_n + 1):
            ws.column_dimensions[get_column_letter(2 + k)].width = 20
        ws.column_dimensions[get_column_letter(3 + weeks_n)].width = 36
        ws.column_dimensions[get_column_letter(4 + weeks_n)].width = 8

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
